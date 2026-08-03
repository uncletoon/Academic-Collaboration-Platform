from __future__ import annotations

import math
import shutil
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(
    r"D:\Toon\My Doc\Classmate\Divine\Final Academic Collaboration Research proposal.docx"
)
OUTPUT = Path(
    r"D:\Toon\My Doc\Classmate\Divine\Final Academic Collaboration Research Book - Chapters 4 and 5.docx"
)
WORK = Path(r"D:\Toon\My Doc\Classmate\Divine\Collaboration\.codex-work")
FIGURES = WORK / "figures"

NAVY = "#17365D"
BLUE = "#2F5FB3"
PALE = "#F7FAFE"
LINE = "#7F8FA6"
BLACK = "#111111"
_BULLET_NUM_ID = None


def add_box(ax, xy, width, height, text, fontsize=10, edge=NAVY, lw=1.5):
    x, y = xy
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.015,rounding_size=0.015",
        linewidth=lw,
        edgecolor=edge,
        facecolor="white",
    )
    ax.add_patch(patch)
    ax.text(
        x + width / 2,
        y + height / 2,
        "\n".join(textwrap.wrap(text, width=max(15, int(width * 62)))),
        ha="center",
        va="center",
        fontsize=fontsize,
        color=BLACK,
        weight="semibold",
    )
    return patch


def arrow(ax, start, end, label="", fontsize=8.5, offset=(0, 0.018), style="-|>"):
    arr = FancyArrowPatch(
        start,
        end,
        arrowstyle=style,
        mutation_scale=12,
        linewidth=1.25,
        color=NAVY,
        connectionstyle="arc3",
    )
    ax.add_patch(arr)
    if label:
        mx = (start[0] + end[0]) / 2 + offset[0]
        my = (start[1] + end[1]) / 2 + offset[1]
        ax.text(
            mx,
            my,
            "\n".join(textwrap.wrap(label, width=28)),
            ha="center",
            va="center",
            fontsize=fontsize,
            color=BLACK,
            bbox=dict(facecolor="white", edgecolor="none", pad=1.2),
        )


def setup_ax(title):
    fig, ax = plt.subplots(figsize=(12, 7.2), dpi=190)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.965, title, ha="center", va="top", fontsize=16, weight="bold", color=NAVY)
    return fig, ax


def save_context_dfd(path):
    fig, ax = setup_ax("Context Data Flow Diagram (Level 0)")
    add_box(ax, (0.38, 0.38), 0.24, 0.22, "Academic Collaboration and Real-Time Notification System", 11)
    add_box(ax, (0.04, 0.68), 0.22, 0.13, "Students, Lecturers and Researchers", 10)
    add_box(ax, (0.74, 0.68), 0.22, 0.13, "Institution and System Administrators", 10)
    add_box(ax, (0.04, 0.12), 0.22, 0.13, "Institutional Data Sources", 10)
    add_box(ax, (0.74, 0.12), 0.22, 0.13, "Email / Browser Clients", 10)
    arrow(ax, (0.26, 0.72), (0.38, 0.56), "registration, posts, messages,\nproject and event actions")
    arrow(ax, (0.38, 0.45), (0.26, 0.20), "feeds, search results,\nstatus and alerts", offset=(-0.01, 0))
    arrow(ax, (0.74, 0.72), (0.62, 0.56), "management rules,\ncontent and reports")
    arrow(ax, (0.62, 0.44), (0.74, 0.20), "system status,\nmoderation results", offset=(0.01, 0))
    arrow(ax, (0.26, 0.19), (0.38, 0.42), "institution and department metadata", offset=(-0.015, 0))
    arrow(ax, (0.62, 0.42), (0.74, 0.19), "real-time notification and message delivery", offset=(0.02, 0))
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_level1_dfd(path):
    fig, ax = setup_ax("Data Flow Diagram (Level 1)")
    # external actors
    add_box(ax, (0.02, 0.66), 0.16, 0.11, "Academic User", 10)
    add_box(ax, (0.82, 0.66), 0.16, 0.11, "Administrator", 10)
    # processes
    processes = [
        (0.23, 0.72, "1.0\nAuthenticate and Manage Profile"),
        (0.43, 0.72, "2.0\nManage Communities and Projects"),
        (0.63, 0.72, "3.0\nManage Events and News"),
        (0.23, 0.39, "4.0\nChat and Notifications"),
        (0.43, 0.39, "5.0\nSearch and Discovery"),
        (0.63, 0.39, "6.0\nAdministration and Reporting"),
    ]
    for x, y, t in processes:
        add_box(ax, (x, y), 0.15, 0.12, t, 8.8)
    # data stores
    stores = [
        (0.12, 0.10, "D1 Identity and Institution Data"),
        (0.34, 0.10, "D2 Community and Project Data"),
        (0.56, 0.10, "D3 Event, News and File Data"),
        (0.78, 0.10, "D4 Message and Notification Data"),
    ]
    for x, y, t in stores:
        add_box(ax, (x, y), 0.17, 0.10, t, 8.5, edge=LINE, lw=1.2)
    arrow(ax, (0.18, 0.71), (0.23, 0.78), "credentials")
    arrow(ax, (0.18, 0.69), (0.43, 0.76), "collaboration actions", offset=(0, -0.02))
    arrow(ax, (0.18, 0.68), (0.63, 0.75), "event actions", offset=(0, 0.025))
    arrow(ax, (0.18, 0.66), (0.23, 0.45), "messages")
    arrow(ax, (0.18, 0.67), (0.43, 0.45), "queries", offset=(0, -0.02))
    arrow(ax, (0.82, 0.71), (0.78, 0.75), "moderation")
    arrow(ax, (0.82, 0.69), (0.78, 0.45), "account and content control")
    for px, py, _ in processes:
        target = min(stores, key=lambda s: abs((s[0] + 0.085) - (px + 0.075)))
        arrow(ax, (px + 0.075, py), (target[0] + 0.085, target[1] + 0.10), "", offset=(0, 0))
    ax.text(
        0.5,
        0.03,
        "All protected flows pass through JWT authentication and active-account checks. "
        "Socket.IO delivers real-time messages and notifications.",
        ha="center",
        va="bottom",
        fontsize=9.5,
        color=BLACK,
    )
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_sequence(path):
    fig, ax = setup_ax("System Sequence Diagram: Event Publication and Real-Time Notification")
    actors = [
        (0.08, "Administrator"),
        (0.28, "React Client"),
        (0.50, "Express API"),
        (0.70, "PostgreSQL"),
        (0.90, "Socket.IO / User Client"),
    ]
    for x, label in actors:
        add_box(ax, (x - 0.075, 0.82), 0.15, 0.09, label, 9)
        ax.plot([x, x], [0.15, 0.82], linestyle="--", linewidth=1, color=LINE)
    messages = [
        (0.08, 0.28, 0.76, "Submit event details"),
        (0.28, 0.50, 0.68, "POST /api/events + JWT"),
        (0.50, 0.70, 0.60, "Validate and INSERT event"),
        (0.70, 0.50, 0.52, "Return saved event"),
        (0.50, 0.70, 0.44, "INSERT notification records"),
        (0.50, 0.90, 0.36, "Emit notification to user room"),
        (0.90, 0.28, 0.28, "Display alert / refresh event feed"),
        (0.28, 0.08, 0.20, "Show publication success"),
    ]
    for sx, ex, y, label in messages:
        direction = "-|>"
        arr = FancyArrowPatch(
            (sx, y),
            (ex, y),
            arrowstyle=direction,
            mutation_scale=11,
            linewidth=1.2,
            color=NAVY,
        )
        ax.add_patch(arr)
        ax.text(
            (sx + ex) / 2,
            y + 0.018,
            label,
            ha="center",
            va="bottom",
            fontsize=8.2,
            color=BLACK,
            bbox=dict(facecolor="white", edgecolor="none", pad=0.8),
        )
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_physical_model(path):
    fig, ax = setup_ax("Physical Data Model")
    groups = [
        (0.03, 0.62, 0.29, 0.23, "Identity", ["institutions", "departments", "users"]),
        (0.355, 0.62, 0.29, 0.23, "Academic Community", ["academic_communities", "community_members", "invitations", "posts / comments / likes"]),
        (0.68, 0.62, 0.29, 0.23, "Projects", ["projects", "project_members", "join_requests", "project_files"]),
        (0.03, 0.25, 0.29, 0.23, "Events and News", ["events", "event_registrations", "news", "news_documents"]),
        (0.355, 0.25, 0.29, 0.23, "Communication", ["chat_rooms", "chat_members", "chat_messages", "notifications"]),
        (0.68, 0.25, 0.29, 0.23, "Physical Platform", ["PostgreSQL relational database", "server-side upload folders", "indexed PK/FK relationships"]),
    ]
    for x, y, w, h, title, items in groups:
        box = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.015,rounding_size=0.015",
            linewidth=1.4,
            edgecolor=NAVY,
            facecolor="white",
        )
        ax.add_patch(box)
        ax.text(x + 0.02, y + h - 0.04, title, ha="left", va="top", fontsize=11, weight="bold", color=NAVY)
        for i, item in enumerate(items):
            ax.text(x + 0.03, y + h - 0.085 - i * 0.035, f"• {item}", ha="left", va="top", fontsize=8.8, color=BLACK)
    arrow(ax, (0.175, 0.62), (0.50, 0.48), "user and institution foreign keys", offset=(0, 0.01))
    arrow(ax, (0.50, 0.62), (0.825, 0.48), "ownership and membership", offset=(0, 0.01))
    arrow(ax, (0.825, 0.62), (0.825, 0.48), "project files and access", offset=(0, 0.015))
    arrow(ax, (0.175, 0.25), (0.50, 0.25), "event and content alerts", offset=(0, 0.02))
    arrow(ax, (0.645, 0.36), (0.68, 0.36), "stored by", offset=(0, 0.02))
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_erd(path):
    fig, ax = setup_ax("Entity Relationship Diagram (Core Relationships)")
    entities = {
        "INSTITUTION": (0.04, 0.72),
        "DEPARTMENT": (0.04, 0.46),
        "USER": (0.32, 0.59),
        "COMMUNITY": (0.60, 0.76),
        "PROJECT": (0.60, 0.50),
        "EVENT": (0.60, 0.24),
        "CHAT_ROOM": (0.32, 0.20),
        "NOTIFICATION": (0.82, 0.50),
    }
    subtitles = {
        "INSTITUTION": "PK id",
        "DEPARTMENT": "PK id · FK institution_id",
        "USER": "PK id · FK institution_id, department_id",
        "COMMUNITY": "PK id · FK created_by, institution_id",
        "PROJECT": "PK id · FK created_by, institution_id",
        "EVENT": "PK id · FK organizer_id, institution_id",
        "CHAT_ROOM": "PK id",
        "NOTIFICATION": "PK id · FK user_id",
    }
    centers = {}
    for name, (x, y) in entities.items():
        w, h = 0.16, 0.115
        add_box(ax, (x, y), w, h, f"{name}\n{subtitles[name]}", 8.7)
        centers[name] = (x + w / 2, y + h / 2)

    rels = [
        ("INSTITUTION", "DEPARTMENT", "1 : many"),
        ("INSTITUTION", "USER", "1 : many"),
        ("DEPARTMENT", "USER", "1 : many"),
        ("USER", "COMMUNITY", "creates / joins"),
        ("USER", "PROJECT", "creates / joins"),
        ("USER", "EVENT", "organizes / registers"),
        ("USER", "CHAT_ROOM", "many : many"),
        ("USER", "NOTIFICATION", "1 : many"),
        ("PROJECT", "NOTIFICATION", "generates"),
        ("EVENT", "NOTIFICATION", "generates"),
    ]
    for a, b, label in rels:
        arrow(ax, centers[a], centers[b], label, fontsize=7.5, offset=(0, 0.012), style="-")
    ax.text(
        0.5,
        0.06,
        "Intersection tables resolve many-to-many relationships: community_members, project_members, "
        "event_registrations and chat_members. Dependent content tables use foreign keys with controlled deletion.",
        ha="center",
        va="center",
        fontsize=9,
        color=BLACK,
        wrap=True,
    )
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def _font(size, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _canvas(title):
    image = Image.new("RGB", (2200, 1320), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(46, True)
    box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((2200 - (box[2] - box[0])) / 2, 38), title, fill=NAVY, font=title_font)
    return image, draw


def _wrap_by_pixels(draw, text, font, max_width):
    output = []
    for original_line in str(text).splitlines() or [""]:
        words = original_line.split()
        if not words:
            output.append("")
            continue
        line = words[0]
        for word in words[1:]:
            trial = f"{line} {word}"
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
                line = trial
            else:
                output.append(line)
                line = word
        output.append(line)
    return "\n".join(output)


def _box(draw, x, y, w, h, text, size=29, edge=NAVY):
    draw.rounded_rectangle((x, y, x + w, y + h), radius=20, fill="white", outline=edge, width=4)
    font = _font(size, True)
    wrapped = _wrap_by_pixels(draw, text, font, w - 44)
    bbox = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=8, align="center")
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(
        (x + (w - tw) / 2, y + (h - th) / 2),
        wrapped,
        fill=BLACK,
        font=font,
        spacing=8,
        align="center",
    )


def _arrow(draw, start, end, label="", size=23, arrowhead=True, label_offset=(0, -10)):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=NAVY, width=4)
    if arrowhead:
        angle = math.atan2(y2 - y1, x2 - x1)
        length = 22
        spread = 0.55
        p1 = (x2 - length * math.cos(angle - spread), y2 - length * math.sin(angle - spread))
        p2 = (x2 - length * math.cos(angle + spread), y2 - length * math.sin(angle + spread))
        draw.polygon([(x2, y2), p1, p2], fill=NAVY)
    if label:
        font = _font(size, False)
        wrapped = _wrap_by_pixels(draw, label, font, 390)
        bbox = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=5, align="center")
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        mx = (x1 + x2) / 2 - tw / 2 + label_offset[0]
        my = (y1 + y2) / 2 - th / 2 + label_offset[1]
        draw.rectangle((mx - 7, my - 4, mx + tw + 7, my + th + 4), fill="white")
        draw.multiline_text((mx, my), wrapped, fill=BLACK, font=font, spacing=5, align="center")


def save_context_dfd_pil(path):
    image, draw = _canvas("Context Data Flow Diagram (Level 0)")
    _box(draw, 825, 485, 550, 300, "Academic Collaboration and Real-Time Notification System", 34)
    _box(draw, 70, 225, 475, 175, "Students, Lecturers and Researchers", 30)
    _box(draw, 1655, 225, 475, 175, "Institution and System Administrators", 30)
    _box(draw, 70, 915, 475, 175, "Institutional Data Sources", 30)
    _box(draw, 1655, 915, 475, 175, "Email and Browser Clients", 30)
    _arrow(draw, (545, 310), (825, 520), "registration, posts, messages,\nproject and event actions", 22)
    _arrow(draw, (825, 565), (545, 355), "feeds, search results,\nstatus and alerts", 22, label_offset=(-20, 30))
    _arrow(draw, (1655, 310), (1375, 520), "management rules\nand content", 22)
    _arrow(draw, (1375, 565), (1655, 355), "reports and\nmoderation results", 22, label_offset=(20, 30))
    _arrow(draw, (545, 1000), (825, 690), "institution and department metadata", 22, label_offset=(-30, 25))
    _arrow(draw, (1375, 690), (1655, 1000), "real-time notifications\nand messages", 22, label_offset=(35, 25))
    image.save(path)


def save_level1_dfd_pil(path):
    image, draw = _canvas("Data Flow Diagram (Level 1)")
    _box(draw, 35, 520, 330, 150, "Academic User", 32)
    _box(draw, 1835, 520, 330, 150, "Administrator", 32)
    processes = [
        (350, 175, "1.0\nAuthenticate and Manage Profile"),
        (850, 125, "2.0\nManage Communities and Projects"),
        (1430, 175, "3.0\nManage Events and News"),
        (350, 870, "4.0\nChat and Notifications"),
        (850, 920, "5.0\nSearch and Discovery"),
        (1430, 870, "6.0\nAdministration and Reporting"),
    ]
    for x, y, text in processes:
        _box(draw, x, y, 420, 175, text, 25)
    _box(
        draw,
        825,
        500,
        550,
        270,
        "POSTGRESQL DATA STORES\n\nD1 Identity and Institution\nD2 Community and Project\nD3 Event, News and Files\nD4 Messages and Notifications",
        23,
        edge=LINE,
    )
    _arrow(draw, (365, 575), (350, 310), "credentials and profile data", 18, label_offset=(-50, -5))
    _arrow(draw, (365, 620), (350, 900), "messages and notification actions", 18, label_offset=(-45, 10))
    _arrow(draw, (1835, 575), (1850, 310), "event and content management", 18, label_offset=(45, -5))
    _arrow(draw, (1835, 620), (1850, 900), "account and reporting control", 18, label_offset=(45, 10))
    process_edges = [
        ((560, 350), (900, 500)),
        ((1060, 300), (1060, 500)),
        ((1430, 350), (1300, 500)),
        ((560, 870), (900, 770)),
        ((1060, 920), (1060, 770)),
        ((1430, 870), (1300, 770)),
    ]
    for start, end in process_edges:
        _arrow(draw, start, end, "", arrowhead=True)
    note = (
        "Protected processes use JWT authentication and active-account checks. "
        "Socket.IO provides real-time message and notification delivery."
    )
    font = _font(25)
    wrapped = _wrap_by_pixels(draw, note, font, 1750)
    draw.multiline_text((225, 1215), wrapped, fill=BLACK, font=font, spacing=6, align="center")
    image.save(path)


def save_sequence_pil(path):
    image, draw = _canvas("System Sequence Diagram: Event Publication and Real-Time Notification")
    actors = [
        (180, "Administrator"),
        (620, "React Client"),
        (1060, "Express API"),
        (1500, "PostgreSQL"),
        (1940, "Socket.IO / User Client"),
    ]
    for x, label in actors:
        _box(draw, x - 160, 155, 320, 125, label, 27)
        y = 300
        while y < 1160:
            draw.line((x, y, x, min(y + 24, 1160)), fill=LINE, width=3)
            y += 42
    messages = [
        (180, 620, 350, "Submit event details"),
        (620, 1060, 455, "POST /api/events + JWT"),
        (1060, 1500, 560, "Validate and INSERT event"),
        (1500, 1060, 665, "Return saved event"),
        (1060, 1500, 770, "INSERT notification records"),
        (1060, 1940, 875, "Emit notification to user room"),
        (1940, 620, 980, "Display alert and refresh feed"),
        (620, 180, 1085, "Show publication success"),
    ]
    for sx, ex, y, label in messages:
        _arrow(draw, (sx, y), (ex, y), label, 22, label_offset=(0, -30))
    image.save(path)


def save_physical_model_pil(path):
    image, draw = _canvas("Physical Data Model")
    groups = [
        (60, 180, "Identity", ["institutions", "departments", "users"]),
        (760, 180, "Academic Community", ["academic_communities", "community_members", "invitations", "posts / comments / likes"]),
        (1460, 180, "Projects", ["projects", "project_members", "join_requests", "project_files"]),
        (60, 720, "Events and News", ["events", "event_registrations", "news", "news_documents"]),
        (760, 720, "Communication", ["chat_rooms", "chat_members", "chat_messages", "notifications"]),
        (1460, 720, "Physical Platform", ["PostgreSQL database", "server upload folders", "PK/FK constraints"]),
    ]
    for x, y, title, items in groups:
        draw.rounded_rectangle((x, y, x + 620, y + 350), radius=22, fill="white", outline=NAVY, width=4)
        draw.text((x + 35, y + 30), title, fill=NAVY, font=_font(34, True))
        for i, item in enumerate(items):
            draw.text((x + 45, y + 100 + i * 55), f"• {item}", fill=BLACK, font=_font(27))
    _arrow(draw, (680, 355), (760, 355), "identity links", 20, label_offset=(0, -34))
    _arrow(draw, (1380, 355), (1460, 355), "ownership", 20, label_offset=(0, -34))
    _arrow(draw, (370, 720), (1070, 530), "event and content alerts", 21)
    _arrow(draw, (1380, 895), (1460, 895), "stored by", 20, label_offset=(0, -34))
    image.save(path)


def save_erd_pil(path):
    image, draw = _canvas("Entity Relationship Diagram (Core Relationships)")
    entities = {
        "INSTITUTION": (60, 165, "PK id"),
        "DEPARTMENT": (60, 560, "PK id · FK institution_id"),
        "USER": (650, 365, "PK id · FK institution_id, department_id"),
        "NOTIFICATION": (650, 110, "PK id · FK user_id"),
        "CHAT_ROOM": (650, 835, "PK id"),
        "COMMUNITY": (1250, 145, "PK id · FK created_by, institution_id"),
        "PROJECT": (1250, 480, "PK id · FK created_by, institution_id"),
        "EVENT": (1250, 815, "PK id · FK organizer_id, institution_id"),
    }
    for name, (x, y, detail) in entities.items():
        _box(draw, x, y, 360, 155, f"{name}\n{detail}", 24)
    _arrow(draw, (240, 320), (240, 560), "1 : many", 18, arrowhead=False, label_offset=(20, 0))
    _arrow(draw, (420, 240), (650, 405), "1 : many", 18, arrowhead=False, label_offset=(0, -18))
    _arrow(draw, (420, 635), (650, 480), "1 : many", 18, arrowhead=False, label_offset=(0, 18))
    _arrow(draw, (1010, 405), (1250, 240), "creates / joins", 18, arrowhead=False, label_offset=(0, -18))
    _arrow(draw, (1010, 442), (1250, 558), "creates / joins", 18, arrowhead=False, label_offset=(0, -18))
    _arrow(draw, (1010, 480), (1250, 885), "organizes / registers", 18, arrowhead=False, label_offset=(15, 0))
    _arrow(draw, (830, 520), (830, 835), "many : many", 18, arrowhead=False, label_offset=(20, 0))
    _arrow(draw, (830, 365), (830, 265), "1 : many", 18, arrowhead=False, label_offset=(20, 0))
    note = (
        "Intersection tables resolve many-to-many relationships: community_members, project_members, "
        "event_registrations and chat_members. Dependent content uses controlled foreign keys."
    )
    font = _font(23)
    wrapped = _wrap_by_pixels(draw, note, font, 1800)
    draw.multiline_text((200, 1190), wrapped, fill=BLACK, font=font, spacing=6, align="center")
    image.save(path)


def create_figures():
    FIGURES.mkdir(parents=True, exist_ok=True)
    save_context_dfd_pil(FIGURES / "dfd-level-0.png")
    save_level1_dfd_pil(FIGURES / "dfd-level-1.png")
    save_sequence_pil(FIGURES / "sequence-event-notification.png")
    save_physical_model_pil(FIGURES / "physical-data-model.png")
    save_erd_pil(FIGURES / "erd-core.png")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_run(run, name="Times New Roman", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        format_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead) :])
        format_run(r2)
    else:
        r = p.add_run(text)
        format_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    if _BULLET_NUM_ID is not None:
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(_BULLET_NUM_ID))
        num_pr.append(ilvl)
        num_pr.append(num_id)
        p_pr.append(num_pr)
    r = p.add_run(text)
    format_run(r)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    for r in p.runs:
        format_run(r, size=16 if level == 1 else (14 if level == 2 else 12), bold=True)
    return p


def add_table(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    for r in cap.runs:
        format_run(r, size=10, italic=True, color=(31, 73, 125))

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_width(cell, widths[j])
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        format_run(r, size=9, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            format_run(r, size=8.5)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_figure(doc, image_path, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    for r in cap.runs:
        format_run(r, size=10, italic=True, color=(31, 73, 125))
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F5F5")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2F5FB3")
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run("\n".join(lines))
    format_run(r, name="Consolas", size=8.0)
    return p


def ensure_list_bullet_style(doc):
    try:
        style = doc.styles["List Bullet"]
    except KeyError:
        style = doc.styles.add_style("List Bullet", 1)
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def install_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 0
    num_id_value = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:hint"), "default")
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id_value))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id_value


def build_document():
    create_figures()
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    ensure_list_bullet_style(doc)
    global _BULLET_NUM_ID
    _BULLET_NUM_ID = install_bullet_numbering(doc)

    reference_paragraph = None
    for p in doc.paragraphs:
        if p.text.strip().upper() == "REFERENCES":
            reference_paragraph = p
            break
    if reference_paragraph is None:
        raise RuntimeError("Could not locate the REFERENCES heading.")

    original_body_count = len(doc._element.body)

    p_break = doc.add_paragraph()
    p_break.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "CHAPTER 4: SYSTEM ANALYSIS, DESIGN, IMPLEMENTATION AND EVALUATION", 1)
    add_heading(doc, "4.0 Introduction", 2)
    add_body(
        doc,
        "This chapter presents the analysis, design, implementation and evaluation of the Cross-Institutional "
        "Academic Collaboration and Real-Time Event Notification System. It connects the study problem and "
        "objectives to the system that was implemented. The chapter explains the available research evidence, "
        "the existing information-sharing situation, the functions of the developed platform, its database and "
        "process design, the technologies used, the visible interfaces and the tests performed.",
    )
    add_body(
        doc,
        "The chapter uses two evidence sources. First, it uses the barriers and user expectations already recorded "
        "in Chapter 3. Second, it uses direct inspection of the supplied source code, database schema, running "
        "application and build output. No questionnaire dataset, completed interview transcripts or signed user "
        "acceptance forms were supplied with the manuscript. Therefore, no response rate, percentage, mean, "
        "Cronbach's alpha value or participant success rate is invented in this chapter. Places that require "
        "verified field statistics are clearly identified for completion after the original records are available.",
    )

    add_heading(doc, "4.1 Data Analysis and Presentation", 2)
    add_body(
        doc,
        "The study planned to collect information through interviews, questionnaires, observation and prototype "
        "testing. The available manuscript identifies repeated qualitative themes, while the implementation "
        "provides objective evidence about the functions that were built. Table 2 separates verified evidence from "
        "information that is still required for a final empirical report.",
    )
    add_table(
        doc,
        "Table 2: Evidence Available for Chapter 4 Analysis",
        ["Evidence item", "Available evidence", "Status and use in this chapter"],
        [
            (
                "Qualitative themes",
                "Chapter 3 records fragmented announcements, difficult repository discovery, institutional barriers, security concerns and demand for immediate notifications.",
                "Available; used for thematic analysis and requirements.",
            ),
            (
                "Population and planned sample",
                "Target population of 10,025 and planned sample of 385 respondents.",
                "Available as a research plan; not treated as completed responses.",
            ),
            (
                "Questionnaire results",
                "Raw responses, completed response count and coded Likert-scale data were not supplied.",
                "Pending; percentages, means and reliability statistics must be inserted from the verified dataset.",
            ),
            (
                "Interview and observation records",
                "Themes are summarized in Chapter 3, but transcripts and observation sheets were not supplied.",
                "Themes reported without fabricated quotations or participant counts.",
            ),
            (
                "Prototype evidence",
                "Source code, PostgreSQL schema, running interface, API health response and production build output.",
                "Available; used for system description and technical evaluation.",
            ),
            (
                "User acceptance evidence",
                "No signed UAT form or participant-level result sheet was supplied.",
                "Acceptance criteria are provided; outcome remains pending verification.",
            ),
        ],
        [1.35, 2.65, 2.35],
    )
    add_body(
        doc,
        "The qualitative information was organized into themes. Each theme was then translated into a design "
        "response so that the analysis directly informed the system rather than remaining separate from the "
        "implementation.",
    )
    add_table(
        doc,
        "Table 3: Thematic Analysis and Design Response",
        ["Theme", "Meaning of the finding", "System response"],
        [
            (
                "Fragmented communication",
                "Academic notices are spread across websites, social media, email and physical channels.",
                "One portal combines news, events, communities, projects and notifications.",
            ),
            (
                "Pull-based discovery",
                "Users must repeatedly visit institutional pages or repositories to find new information.",
                "Event feeds, global search and stored notifications make information easier to discover.",
            ),
            (
                "Institutional isolation",
                "Students and staff have limited visibility of people and work outside their institution.",
                "Institution-linked profiles, cross-institution communities and eligible project discovery support connection.",
            ),
            (
                "Coordination difficulty",
                "Project membership, shared files, roles and discussions are managed through separate tools.",
                "Private project workspaces provide roles, join requests, members and supportive documents.",
            ),
            (
                "Need for immediate feedback",
                "Delayed messages and event notices reduce timely participation.",
                "Socket.IO rooms deliver live chat messages and user-specific notifications.",
            ),
            (
                "Trust, privacy and security",
                "Users and institutions are concerned about access, ownership and misuse of academic content.",
                "JWT authentication, active-account checks, role controls, private projects and controlled file routes protect access.",
            ),
        ],
        [1.35, 2.45, 2.55],
    )
    add_body(
        doc,
        "The quantitative analysis planned in Chapter 3 cannot be completed responsibly without the original "
        "response file. When it becomes available, the final version should report the completed sample, response "
        "rate, participant characteristics, item frequencies, means, standard deviations, reliability results and "
        "pre-test/post-test or task-completion comparisons. These values should be calculated from the raw data and "
        "should not be estimated from the planned sample.",
    )
    add_table(
        doc,
        "Table 4: Field Statistics to Complete from the Verified Dataset",
        ["Required statistic", "Value to enter", "Source required"],
        [
            ("Completed questionnaires", "[Insert verified number]", "Final response dataset"),
            ("Response rate", "[Insert verified percentage]", "Completed responses divided by 385"),
            ("Interview participants", "[Insert verified number by role]", "Interview register and consent records"),
            ("Usability participants", "[Insert verified number by role]", "Prototype test attendance sheet"),
            ("Task completion and time", "[Insert verified values]", "Observation and usability task sheets"),
            ("System usability / satisfaction", "[Insert verified mean or percentage]", "Prototype evaluation forms"),
            ("Cronbach's alpha", "[Insert calculated coefficient]", "Final Likert-scale item dataset"),
        ],
        [2.1, 1.65, 2.6],
    )

    add_heading(doc, "4.2 Interpretation and Discussion of Findings", 2)
    add_body(
        doc,
        "The findings show that the main problem is not only lack of internet access. It is also the absence of one "
        "application layer that connects people, projects, events and communication across institutions. Better "
        "network infrastructure cannot by itself make a student aware of a workshop, help a lecturer find a partner "
        "or give a project team a shared private workspace. The developed system responds by bringing these actions "
        "into one authenticated portal.",
    )
    add_body(
        doc,
        "The theme of fragmented communication supports the need for a unified feed and global search. In the "
        "developed platform, users can search across people, communities, projects, events and news. This reduces "
        "the need to check many separate channels. The events module also supports publication, filtering, capacity "
        "management and registration, which creates a direct path from discovery to participation.",
    )
    add_body(
        doc,
        "The theme of institutional isolation is addressed through profiles linked to institutions and departments, "
        "cross-institutional communities and collaboration projects. Communities provide posts, comments, likes, "
        "invitations and membership rules. Projects add requirements, access scope, join requests, team roles and "
        "supportive files. These functions change collaboration from an informal announcement into a managed "
        "academic workflow.",
    )
    add_body(
        doc,
        "The need for immediate communication is addressed with Socket.IO. Messages are saved in PostgreSQL and "
        "broadcast to chat rooms, while notifications are stored and emitted to a room belonging to the intended "
        "user. This gives both a permanent history and immediate delivery. The implemented prototype does not "
        "contain an MQTT broker. MQTT was proposed in the earlier design, but the current code uses Socket.IO over "
        "web-compatible real-time transport. This difference should be treated as an implementation decision and a "
        "future extension, not as a completed MQTT feature.",
    )
    add_body(
        doc,
        "Security findings are reflected in authentication and authorization controls. Passwords are hashed, "
        "sessions use signed JWT tokens, protected routes check both the token and the current account status, and "
        "administrator functions are restricted by role. Projects remain private, with controlled eligibility and "
        "owner approval. These controls answer important trust concerns, although production deployment still "
        "requires stronger secret management, secure cookies or protected token storage, file scanning, HTTPS, "
        "auditing and regular security testing.",
    )
    add_body(
        doc,
        "The running system demonstrates technical feasibility, but it does not prove the size of the platform's "
        "effect on communication or event participation. Such impact requires verified field data and comparison "
        "over time. The correct interpretation is therefore that the system provides the mechanisms needed to "
        "improve collaboration and awareness; the degree of improvement must be established through the planned "
        "multi-user evaluation.",
    )

    add_heading(doc, "4.3 Summary of Findings and System Requirements", 2)
    add_body(
        doc,
        "The analysis produced a clear set of system requirements. The most important requirement was to combine "
        "identity, discovery, collaboration, events and real-time communication in one controlled environment.",
    )
    add_table(
        doc,
        "Table 5: Summary of Functional System Requirements",
        ["ID", "Requirement", "Priority", "Implemented evidence"],
        [
            ("FR-01", "Register, sign in and maintain an institution-linked profile.", "High", "Authentication and profile routes and pages"),
            ("FR-02", "Create, discover, join and manage academic communities.", "High", "Community membership, invitations and privacy rules"),
            ("FR-03", "Publish posts and support comments and likes inside communities.", "Medium", "Post, comment and like controllers"),
            ("FR-04", "Create private collaboration projects with requirements and access scope.", "High", "Projects and project join-request workflow"),
            ("FR-05", "Manage project members, roles and supportive documents.", "High", "Membership and protected file functions"),
            ("FR-06", "Publish, update, discover and register for academic events.", "High", "Event routes and event workspace"),
            ("FR-07", "Send direct and group messages in real time.", "High", "Chat rooms, stored messages and Socket.IO broadcast"),
            ("FR-08", "Create, store, deliver and mark notifications as read.", "High", "Notification service, table and user rooms"),
            ("FR-09", "Search across the main academic content types.", "Medium", "Global search endpoint and search interface"),
            ("FR-10", "Administer users, roles, communities, events and summary statistics.", "High", "Role-protected administration functions"),
        ],
        [0.55, 3.15, 0.7, 1.95],
    )

    add_heading(doc, "4.4 Description of the Existing Academic Collaboration System", 2)
    add_body(
        doc,
        "Before the proposed platform, academic collaboration and event communication were handled through a "
        "collection of separate methods. Institutional websites published selected notices, repositories stored "
        "research outputs, social media carried informal announcements, email reached limited mailing lists and "
        "physical notice boards served people already on campus. These methods were useful individually, but they "
        "did not form one connected workflow.",
    )
    add_body(
        doc,
        "The existing approach was mainly pull-based. A user had to know which website, page or social account to "
        "visit. Information could be duplicated, delayed or missed. There was no shared identity layer for academic "
        "users from different institutions, no consistent project membership process, no protected common file "
        "space and no single event-registration path. Conversations often moved to unrelated messaging tools where "
        "academic context, institutional roles and project history were difficult to preserve.",
    )
    add_body(
        doc,
        "Administration was also fragmented. Institutions could publish information on their own channels, but they "
        "had limited visibility of how users interacted across communities, projects and events. This made it hard "
        "to manage participation, apply common rules or maintain one reliable record of activity. These limitations "
        "formed the basis for the developed system.",
    )

    add_heading(doc, "4.5 Description of the Developed System", 2)
    add_body(
        doc,
        "The developed solution is a web-based academic collaboration portal. It provides one interface for students, "
        "lecturers, researchers and administrators. The frontend is a React single-page application. The backend is "
        "an Express API connected to PostgreSQL. Socket.IO provides live chat and notification delivery. The system "
        "uses institution and department records to give academic context to users and to enforce selected access "
        "rules.",
    )
    add_body(
        doc,
        "A user signs in and enters a personalized portal. From the portal, the user can discover communities, "
        "review private collaboration projects, check events, read news, search for academic content, communicate "
        "with other users and receive alerts. Project and community owners manage membership. Administrators can "
        "monitor summary information, manage account status and roles, and remove selected content when necessary.",
    )

    add_heading(doc, "4.5.1 Functional Module Descriptions", 3)
    modules = [
        (
            "Authentication and profile module: ",
            "registers users, validates credentials, hashes passwords, issues JWT tokens, maintains profile details and links users to institutions and departments.",
        ),
        (
            "Institution metadata module: ",
            "provides institution and department lists used during registration, profile management and access decisions.",
        ),
        (
            "Academic community module: ",
            "supports public, private and institution-restricted communities, invitations, membership, posts, comments and likes.",
        ),
        (
            "Collaboration project module: ",
            "creates private projects, defines requirements, controls who may request access, supports owner approval, assigns roles and manages shared documents.",
        ),
        (
            "Event module: ",
            "allows authorized users to create and update events, add physical or online meeting information, manage capacity and let users register or withdraw.",
        ),
        (
            "News module: ",
            "publishes academic news with a feature image, optional external link and protected supporting documents. Content changes are restricted to administrators.",
        ),
        (
            "Chat module: ",
            "supports direct and group rooms, stores message history and broadcasts new messages to active room members.",
        ),
        (
            "Notification module: ",
            "stores user-specific alerts, emits them in real time and lets users mark one or all notifications as read.",
        ),
        (
            "Search module: ",
            "returns matching users, communities, eligible projects, events and news from one query.",
        ),
        (
            "Administration module: ",
            "shows system summary statistics and supports account suspension, role changes and selected content moderation.",
        ),
    ]
    for lead, rest in modules:
        add_body(doc, lead + rest, bold_lead=lead)

    add_heading(doc, "4.5.2 System Configuration", 3)
    add_table(
        doc,
        "Table 6: System Configuration",
        ["Layer", "Implemented configuration", "Minimum deployment requirement"],
        [
            ("Client", "React 18, Vite, Axios/fetch, Tailwind CSS, Lucide icons and Socket.IO client", "Modern Chrome, Edge, Firefox or mobile browser"),
            ("Application server", "Node.js with Express, CORS, JSON/form parsing, Multer and Socket.IO", "Node.js 18 or later; controlled process service"),
            ("Database", "PostgreSQL relational database", "PostgreSQL 14 or later; dedicated database and backup account"),
            ("Authentication", "bcrypt password hashing and JSON Web Tokens", "Strong JWT secret, short expiry policy and HTTPS"),
            ("File storage", "Server upload folders plus database metadata", "Protected directory, file limits, malware scanning and backup"),
            ("Network", "HTTP API and Socket.IO real-time connection", "HTTPS/WSS reverse proxy, stable internet and firewall rules"),
            ("Development", "npm scripts, Vite production build and modular controllers", "Version control, environment-specific configuration and CI"),
        ],
        [1.15, 2.95, 2.25],
    )
    add_body(
        doc,
        "For local development, the frontend is served separately from the API. In production, a reverse proxy should "
        "serve the compiled frontend and forward API and Socket.IO traffic to the Node.js service. Database passwords, "
        "JWT secrets and allowed origins must be supplied through protected environment variables rather than source "
        "code defaults.",
    )

    add_heading(doc, "4.5.3 Non-Functional Requirements", 3)
    add_table(
        doc,
        "Table 7: Non-Functional Requirements",
        ["Category", "Requirement"],
        [
            ("Security", "Authenticate protected requests, enforce current account status and roles, hash passwords, validate uploads and use HTTPS/WSS in production."),
            ("Performance", "Load normal pages and search results without unnecessary delay; deliver live messages promptly under expected institutional traffic."),
            ("Availability", "Recover the API and database after service interruption and maintain tested backups."),
            ("Usability", "Use clear labels, predictable navigation, readable text, helpful validation and responsive layouts."),
            ("Accessibility", "Support keyboard use, semantic headings, visible focus, useful alternative text and suitable colour contrast."),
            ("Scalability", "Allow more institutions, users, communities, projects and concurrent sockets without redesigning the basic data model."),
            ("Maintainability", "Keep route, controller, service, configuration and interface concerns modular and documented."),
            ("Data integrity", "Use primary keys, foreign keys, checks, uniqueness rules and transactions for related updates."),
            ("Compatibility", "Operate on common desktop and mobile browsers and standard institutional networks."),
            ("Auditability", "Record security-sensitive administrative and content actions in a future audit log."),
        ],
        [1.4, 4.95],
    )

    add_heading(doc, "4.6 Illustration of the System Design", 2)
    add_body(
        doc,
        "The diagrams in this section use a white background and simple labels so that the system boundary, data "
        "movement and relationships remain clear when printed in black and white.",
    )

    add_heading(doc, "4.6.1 Data Flow Diagrams", 3)
    add_body(
        doc,
        "The Level 0 diagram treats the whole platform as one process. Academic users send account, collaboration, "
        "project, event and message data to the system and receive feeds, results and notifications. Administrators "
        "supply management decisions and receive operational information.",
    )
    add_figure(doc, FIGURES / "dfd-level-0.png", "Figure 4: Data Flow Diagram Level 0 (Context Diagram)", 6.15)
    add_body(
        doc,
        "The Level 1 diagram divides the system into six main processes and four logical data stores. Protected "
        "operations pass through authentication and current-account checking before a controller reads or changes "
        "data. The communication process adds real-time delivery to the stored records.",
    )
    add_figure(doc, FIGURES / "dfd-level-1.png", "Figure 5: Data Flow Diagram Level 1", 6.15)

    add_heading(doc, "4.6.2 Use-Case or System Sequence Diagram", 3)
    add_body(
        doc,
        "The sequence diagram illustrates a common event-notification use case. An administrator publishes an event "
        "through the React interface. The API authenticates and validates the request, writes the event to PostgreSQL "
        "and creates relevant notification records. Socket.IO then sends an immediate alert to the intended user "
        "rooms. The receiving client displays the alert and can refresh the event feed.",
    )
    add_figure(
        doc,
        FIGURES / "sequence-event-notification.png",
        "Figure 6: System Sequence Diagram for Event Publication and Notification",
        6.15,
    )

    add_heading(doc, "4.6.3 Database Normalization", 3)
    add_body(
        doc,
        "The PostgreSQL database is organized to reduce duplication and protect data integrity. In First Normal Form "
        "(1NF), each table has identifiable rows and fields contain one logical value. Repeating relationships are "
        "placed in separate tables. For example, community members are stored in community_members rather than in a "
        "comma-separated field inside academic_communities.",
    )
    add_body(
        doc,
        "In Second Normal Form (2NF), non-key values depend on the full key. Intersection tables such as "
        "community_members, project_members, event_registrations and chat_members use both foreign keys as their "
        "primary key. Attributes such as joined_at or role describe the complete membership relationship rather "
        "than only one side of it.",
    )
    add_body(
        doc,
        "In Third Normal Form (3NF), descriptive information is separated so that non-key fields do not depend on "
        "other non-key fields. Institution names and locations are kept in institutions, department names are kept "
        "in departments, and users store only the related identifiers. Similarly, chat room details, messages and "
        "membership are separated. This prevents update, insertion and deletion anomalies.",
    )
    add_body(
        doc,
        "The projects.requirements field is stored as a JSONB array because it is a bounded document-like property "
        "of one project rather than a relationship reused across the system. If requirements later need independent "
        "owners, scoring or reuse, they should be moved into a project_requirements table.",
    )

    add_heading(doc, "4.6.4 Data Dictionary", 3)
    add_body(
        doc,
        "Table 8 summarizes the implemented relations. PK means primary key and FK means foreign key. The complete "
        "column definitions and constraints are maintained in the supplied schema.sql file.",
    )
    add_table(
        doc,
        "Table 8: Data Dictionary",
        ["Relation", "Key fields", "Important attributes and purpose"],
        [
            ("institutions", "PK id", "name, type, location; stores participating institutions"),
            ("departments", "PK id; FK institution_id", "name; stores departments within an institution"),
            ("users", "PK id; FK institution_id, department_id", "email, password_hash, full_name, role, bio, avatar_url, status"),
            ("academic_communities", "PK id; FK created_by, institution_id", "name, description, category, privacy_type"),
            ("community_members", "PK/FK community_id + user_id", "joined_at; resolves community membership"),
            ("community_invitations", "PK id; FK community_id, user_id", "status, created_at; controls invitations"),
            ("posts", "PK id; FK user_id, community_id", "title, content, created_at"),
            ("comments", "PK id; FK post_id, user_id", "content, created_at"),
            ("likes", "PK/FK post_id + user_id", "created_at; prevents duplicate likes"),
            ("projects", "PK id; FK created_by, institution_id", "title, description, requirements, status, access_scope"),
            ("project_members", "PK/FK project_id + user_id", "role, joined_at"),
            ("project_join_requests", "PK id; FK project_id, user_id", "status, requested_at, responded_at, responded_by"),
            ("project_files", "PK id; FK project_id, uploaded_by", "filename, filepath, mime_type, file_size, uploaded_at"),
            ("events", "PK id; FK organizer_id, institution_id", "title, description, event_date, location, meeting_link, capacity"),
            ("event_registrations", "PK/FK event_id + user_id", "registered_at"),
            ("news", "PK id; FK created_by", "title, description, category, feature_image, external_link"),
            ("news_documents", "PK id; FK news_id", "filename, filepath, mime_type, file_size"),
            ("chat_rooms", "PK id", "name, is_group, created_at"),
            ("chat_members", "PK/FK room_id + user_id", "joined_at"),
            ("chat_messages", "PK id; FK room_id, sender_id", "message, created_at"),
            ("notifications", "PK id; FK user_id", "title, content, type, is_read, link, created_at"),
        ],
        [1.55, 2.1, 2.7],
    )

    add_heading(doc, "4.6.5 Physical Data Model", 3)
    add_body(
        doc,
        "The physical model is implemented in PostgreSQL with SERIAL integer keys, VARCHAR and TEXT fields, "
        "TIMESTAMP WITH TIME ZONE values, BOOLEAN flags and one JSONB array. Foreign keys use controlled delete "
        "actions such as CASCADE or SET NULL. Uploaded files are stored on the server, while their names, paths, "
        "types and sizes are recorded in PostgreSQL.",
    )
    add_figure(doc, FIGURES / "physical-data-model.png", "Figure 7: Physical Data Model", 6.15)

    add_heading(doc, "4.6.6 Entity Relationship Diagram", 3)
    add_body(
        doc,
        "The core entity relationship diagram shows how users connect to academic structures and platform activity. "
        "One institution has many departments and users. A user may create or join many communities and projects, "
        "organize or register for events, join chat rooms and receive many notifications. Intersection relations "
        "resolve many-to-many relationships and dependent tables store posts, files, registrations and messages.",
    )
    add_figure(doc, FIGURES / "erd-core.png", "Figure 8: Entity Relationship Diagram", 6.15)

    add_heading(doc, "4.7 System Implementation and Coding", 2)
    add_heading(doc, "4.7.1 Introduction", 3)
    add_body(
        doc,
        "Implementation converted the requirements into a modular web application. The frontend presents a "
        "responsive interface and calls the API. The backend validates requests, applies authorization rules, "
        "queries PostgreSQL and emits real-time events. The code is divided into configuration, middleware, routes, "
        "controllers and services so that each part has a clear responsibility.",
    )

    add_heading(doc, "4.7.2 Implementation Tools and Technologies", 3)
    add_table(
        doc,
        "Table 9: Implementation Tools and Technologies",
        ["Tool or technology", "Use in the developed system"],
        [
            ("React 18", "Builds reusable user-interface pages and manages client state."),
            ("Vite 5", "Runs the development server and produces the optimized frontend build."),
            ("Tailwind CSS and CSS", "Provide responsive layout, typography, spacing and component styling."),
            ("Lucide React", "Provides consistent interface icons."),
            ("Node.js and Express", "Run the API, middleware, routing and business logic."),
            ("PostgreSQL and pg", "Store structured application records and enforce relationships."),
            ("Socket.IO", "Delivers live chat messages and user-specific notifications."),
            ("bcryptjs", "Hashes and compares account passwords."),
            ("jsonwebtoken", "Creates and verifies JWT authentication tokens."),
            ("Multer", "Validates and stores uploaded avatars, news documents and project files."),
            ("Git-compatible source structure", "Supports controlled development and future continuous integration."),
        ],
        [2.0, 4.35],
    )

    add_heading(doc, "4.7.3 System Screenshots and Selected Source Code", 3)
    add_body(
        doc,
        "Figure 9 shows the authenticated home page. It gives the user one starting point for communities, projects, "
        "events and news. The page also summarizes the academic network and highlights upcoming activity.",
    )
    add_figure(doc, FIGURES / "system-home.png", "Figure 9: Authenticated Collaboration Home Page", 6.35)
    add_body(
        doc,
        "Figure 10 shows the private project directory. It displays eligible projects, membership status, access "
        "scope and project counts. A project owner can create a workspace, define requirements, review join requests "
        "and share documents with accepted members.",
    )
    add_figure(doc, FIGURES / "system-projects.png", "Figure 10: Collaboration Project Workspace", 6.35)
    add_body(
        doc,
        "Figure 11 shows the event workspace. Users can filter upcoming, registered, created and past events. "
        "Authorized creators can publish an event and provide location, capacity and an online meeting link.",
    )
    add_figure(doc, FIGURES / "system-events.png", "Figure 11: Academic Event Workspace", 6.35)
    add_body(
        doc,
        "The following selected code shows how a notification is stored and then delivered to the intended user's "
        "Socket.IO room. Storing the record first ensures that the user can still view the notification later.",
    )
    add_code_block(
        doc,
        [
            "const result = await query(",
            "  `INSERT INTO notifications (user_id, title, content, type, link)",
            "   VALUES ($1, $2, $3, $4, $5) RETURNING *`,",
            "  [userId, title, content, type, link || null]",
            ");",
            "sendNotification(userId, result.rows[0]);",
        ],
    )
    add_body(
        doc,
        "The Socket.IO helper targets a user-specific room. This prevents a normal notification from being broadcast "
        "to every connected client.",
    )
    add_code_block(
        doc,
        [
            "function sendNotification(userId, notification) {",
            "  if (!io) return;",
            "  const room = `user_${String(userId)}`;",
            "  io.to(room).emit('notification', notification);",
            "}",
        ],
    )
    add_body(
        doc,
        "Protected API routes pass through token authentication and a live account-status check. Administrative "
        "routes add a role restriction, as shown below.",
    )
    add_code_block(
        doc,
        [
            "router.use(authenticateToken);",
            "router.use(checkUserActive);",
            "router.get('/admin/stats',",
            "  authorizeRoles('admin'),",
            "  adminCtrl.getAdminStats",
            ");",
        ],
    )

    add_heading(doc, "4.8 System Testing", 2)
    add_heading(doc, "4.8.1 Introduction", 3)
    add_body(
        doc,
        "Testing examined whether the supplied application could be built, whether the backend files were valid, "
        "whether the service responded and whether protected content rejected an unauthenticated request. The "
        "running authenticated interface was also inspected on the home, project and event pages. A dedicated "
        "automated unit-test suite was not included in the supplied project, so this chapter distinguishes verified "
        "technical checks from tests that should be added before production deployment.",
    )

    add_heading(doc, "4.8.2 Testing Objectives", 3)
    for item in [
        "Confirm that the frontend compiles into a production build.",
        "Confirm that backend JavaScript files contain no syntax errors.",
        "Confirm that the application server responds through its health endpoint.",
        "Confirm that protected API routes reject requests without a token.",
        "Confirm that authenticated pages can load data from the API and present the main modules.",
        "Define validation, integration, functional, system and user-acceptance checks for the final release.",
    ]:
        add_bullet(doc, item)

    add_table(
        doc,
        "Table 10: Verified Technical Test Results",
        ["Test ID", "Check performed", "Expected result", "Observed result", "Status"],
        [
            ("VT-01", "Vite production build", "Frontend compiles successfully", "1,555 modules transformed; build completed", "PASS"),
            ("VT-02", "Backend JavaScript syntax", "All source files parse", "18 files checked; 0 failures", "PASS"),
            ("VT-03", "GET /health", "HTTP 200 and UP status", "Server returned UP", "PASS"),
            ("VT-04", "GET /api/projects without JWT", "Request rejected", "HTTP 401 returned", "PASS"),
            ("VT-05", "Authenticated interface inspection", "Home, projects and events render", "All three pages rendered with live records", "PASS"),
        ],
        [0.65, 1.75, 1.55, 1.75, 0.65],
    )

    add_heading(doc, "4.8.3 Unit Testing", 3)
    add_body(
        doc,
        "Unit testing should isolate small functions and controller decisions. The source contains clear units for "
        "project ID parsing, requirement cleaning, JWT verification, role checking, file filtering and notification "
        "delivery. Syntax validation passed, but a repeatable automated unit-test framework was not supplied. The "
        "following cases should be implemented with mocked database and Socket.IO dependencies.",
    )
    add_table(
        doc,
        "Table 11: Unit Test Case Matrix",
        ["Case", "Input or condition", "Expected behavior", "Current evidence"],
        [
            ("UT-01", "Missing email, password, name or role at registration", "Return HTTP 400", "Controller rule present"),
            ("UT-02", "Invalid or expired JWT", "Return HTTP 403", "Middleware rule present"),
            ("UT-03", "Project without title, description or requirement", "Reject creation", "Controller rule present"),
            ("UT-04", "Invalid join-request action", "Accept only accept or reject", "Controller rule present"),
            ("UT-05", "Executable upload type", "Reject unsafe upload", "Multer filter present"),
            ("UT-06", "Notification for one user", "Emit only to user_<id> room", "Socket helper present"),
        ],
        [0.7, 2.0, 2.0, 1.65],
    )

    add_heading(doc, "4.8.4 Validation Testing", 3)
    add_body(
        doc,
        "Validation testing checks that incorrect or incomplete information is stopped before it damages stored data. "
        "The application validates required account fields, unique email addresses, allowed roles, project "
        "requirements, access scope, event values and upload limits. PostgreSQL adds check, unique and foreign-key "
        "constraints. Final validation testing should also cover invalid dates, negative capacities, empty messages, "
        "oversized files, unsupported image types and client-side/server-side error consistency.",
    )

    add_heading(doc, "4.8.5 Integration Testing", 3)
    add_body(
        doc,
        "Integration testing checks the connections between the React client, Express API, PostgreSQL database, file "
        "storage and Socket.IO. The authenticated pages loaded records from the live API, showing that the main "
        "client-API-database path was working during inspection. The project build and backend health checks also "
        "passed. A full repeatable integration suite should create temporary test data, verify database changes, "
        "observe socket delivery and remove the temporary records inside isolated test transactions.",
    )
    add_table(
        doc,
        "Table 12: Integration Test Scenarios",
        ["Scenario", "Integrated components", "Expected result"],
        [
            ("User signs in", "React, API, bcrypt, JWT, PostgreSQL", "Valid user receives a token and profile; invalid credentials are rejected."),
            ("Project join request", "Project page, API, database, notification service, Socket.IO", "Owner receives a stored and live alert; approved user becomes a member."),
            ("Event registration", "Event page, API and registration table", "Registration is added once and capacity information updates."),
            ("Chat message", "Chat page, API, message table and Socket.IO room", "Message is stored and appears for other room members."),
            ("Protected file access", "Project/news page, API authorization, file metadata and storage", "Only eligible authenticated users receive the file."),
        ],
        [1.45, 2.5, 2.4],
    )

    add_heading(doc, "4.8.6 Functional and System Testing", 3)
    add_body(
        doc,
        "Functional testing should follow complete user journeys: registration, profile completion, community "
        "participation, project creation and approval, document sharing, event publication and registration, "
        "messaging, notifications, search and administration. System testing should repeat these journeys under "
        "different roles, browsers, screen sizes and network conditions. It should also measure response time, "
        "concurrent socket behavior, recovery after database interruption, backup restoration and permission "
        "boundaries.",
    )
    add_body(
        doc,
        "The supplied system passed the verified build, syntax, service and authentication-boundary checks shown in "
        "Table 10. No evidence was supplied for high-load, recovery, cross-browser or long-duration testing. These "
        "items should be completed before institutional deployment.",
    )

    add_heading(doc, "4.8.7 User Acceptance Testing", 3)
    add_body(
        doc,
        "User Acceptance Testing (UAT) determines whether intended users consider the system useful and ready for "
        "their academic work. The final UAT should include students, lecturers, researchers and administrators from "
        "more than one institution. Participants should complete the same tasks with clear success criteria and "
        "recorded assistance, time, errors and satisfaction.",
    )
    add_table(
        doc,
        "Table 13: User Acceptance Test Criteria",
        ["Acceptance task", "Success criterion", "Verified result"],
        [
            ("Create or update a profile", "User completes the task without researcher intervention", "[Insert from signed UAT record]"),
            ("Find and join an eligible community", "Correct privacy rule is applied and membership is visible", "[Insert from signed UAT record]"),
            ("Request access to a project", "Owner receives the request and can accept or reject it", "[Insert from signed UAT record]"),
            ("Publish or register for an event", "Event is stored and registration status changes correctly", "[Insert from signed UAT record]"),
            ("Send and receive a message", "Other room member receives the stored message in real time", "[Insert from signed UAT record]"),
            ("Find content using global search", "User locates the intended person or academic item", "[Insert from signed UAT record]"),
            ("Overall acceptance", "Agreed threshold is met for usefulness, ease of use and intention to adopt", "[Insert calculated UAT outcome]"),
        ],
        [2.2, 2.65, 1.5],
    )
    add_body(
        doc,
        "Because no completed UAT forms or participant-level result sheet were provided, the system cannot be "
        "academically described as user-accepted at this stage. The interface and core technical flow are ready for "
        "a controlled pilot, after which the verified results should replace the placeholders in Table 13.",
    )

    add_heading(doc, "4.9 Chapter Summary", 2)
    add_body(
        doc,
        "This chapter translated the documented collaboration problems into functional and non-functional "
        "requirements and described the implemented solution. The system combines user identity, academic "
        "communities, private projects, files, events, news, search, chat, notifications and administration in one "
        "React, Express, PostgreSQL and Socket.IO platform. The database was organized using normalized relations "
        "and controlled foreign keys. Diagrams, screenshots and selected code demonstrated the design and "
        "implementation.",
    )
    add_body(
        doc,
        "Technical verification confirmed that the frontend production build completed, all 18 backend JavaScript "
        "files passed syntax checking, the service health endpoint returned UP, and an unauthenticated protected "
        "request was rejected with HTTP 401. The chapter also identified the remaining evidence gap: verified survey "
        "statistics and completed UAT records are still needed before claims about measured user impact can be made. "
        "MQTT and live document co-editing remain proposed extensions rather than implemented functions.",
    )

    p_break2 = doc.add_paragraph()
    p_break2.add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "CHAPTER 5: CONCLUSIONS AND RECOMMENDATIONS", 1)
    add_heading(doc, "5.0 Introduction", 2)
    add_body(
        doc,
        "This chapter concludes the study by responding directly to the three research objectives and questions. It "
        "also presents recommendations for improving the technology, strengthening institutional ownership, "
        "supporting user adoption, protecting information and deploying the system responsibly. The conclusions "
        "are limited to the evidence available in the manuscript, source code, running prototype and technical "
        "checks.",
    )

    add_heading(doc, "5.1 Conclusion", 2)
    add_body(
        doc,
        "The general objective was to design, implement and evaluate a unified web-based environment for academic "
        "collaboration and real-time event notification. This objective was substantially achieved at prototype "
        "level. A working portal was developed with institution-linked accounts, communities, private projects, "
        "events, news, search, chat, notifications and administration. The prototype passed the recorded build, "
        "syntax, health and authentication-boundary checks. Full empirical evaluation remains incomplete until the "
        "planned survey and UAT records are supplied and analysed.",
    )
    add_body(
        doc,
        "Specific Objective 1 and Research Question 1: The study sought to identify the barriers affecting "
        "cross-institutional collaboration, research visibility and immediate academic communication. The "
        "documented findings identify fragmented communication channels, pull-based discovery, institutional "
        "isolation, weak coordination of joint work, delayed awareness and concerns about privacy and intellectual "
        "property. Therefore, the answer to the first research question is that the problem combines technical "
        "fragmentation, institutional access boundaries and human trust or adoption concerns. It cannot be solved by "
        "network connectivity alone.",
        bold_lead="Specific Objective 1 and Research Question 1: ",
    )
    add_body(
        doc,
        "Specific Objective 2 and Research Question 2: The study sought to implement real-time event distribution "
        "and examine how a cross-institutional platform affects communication and knowledge sharing. The prototype "
        "implements stored notifications and Socket.IO-based push delivery, supported by events, communities, "
        "projects, files and chat. These functions create a practical improvement over separate static channels "
        "because discovery, communication and action occur in one environment. However, the implemented system does "
        "not use MQTT, and the size of the effect on users cannot yet be quantified without the missing field "
        "dataset. The conclusion is therefore one of demonstrated technical capability, not a statistically proven "
        "impact.",
        bold_lead="Specific Objective 2 and Research Question 2: ",
    )
    add_body(
        doc,
        "Specific Objective 3 and Research Question 3: The study sought to develop an architecture for real-time "
        "collaboration, notifications and co-editing, and to determine the effect on event participation. The "
        "implemented architecture supports live messaging and notifications, structured communities, private "
        "project workspaces, file sharing, event discovery and registration. These features provide a direct "
        "mechanism through which awareness and participation may improve. Real-time document co-editing is not "
        "present in the supplied code, and no verified before-and-after participation data were supplied. The "
        "prototype should therefore be described as ready for a controlled institutional pilot, not as proof of a "
        "measured participation increase.",
        bold_lead="Specific Objective 3 and Research Question 3: ",
    )
    add_body(
        doc,
        "Overall, the project demonstrates that a unified, role-aware and real-time academic portal is technically "
        "feasible in the study context. Its main contribution is an integrated foundation that turns separate "
        "announcements and informal conversations into searchable, managed and persistent academic activity. The "
        "next stage is to strengthen security and operations, complete the missing real-time extensions and conduct "
        "a documented multi-institution evaluation.",
    )

    add_heading(doc, "5.2 Recommendations", 2)
    add_heading(doc, "5.2.1 Technical Recommendations", 3)
    for item in [
        "Add an automated test suite for controller units, API integration, database transactions, Socket.IO delivery and end-to-end user journeys.",
        "Implement MQTT only where broker-based topic distribution provides a clear advantage, and document how it works with the existing Socket.IO layer.",
        "Add real-time document co-editing through a proven Operational Transformation or Conflict-free Replicated Data Type service instead of describing file sharing as co-editing.",
        "Add pagination, database indexes, background jobs and a shared Socket.IO adapter before supporting a large number of institutions and concurrent users.",
        "Provide a progressive web application or mobile client with low-bandwidth behavior, cached reading and graceful reconnection.",
        "Improve accessibility through a formal WCAG review, keyboard testing, screen-reader testing and accessible error messages.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.2.2 Institutional Recommendations", 3)
    for item in [
        "Create a governance committee with representatives from participating institutions, students, lecturers, researchers, IT teams and data-protection staff.",
        "Agree on institution onboarding, verified identity, content ownership, moderation, intellectual-property and dispute-resolution policies.",
        "Assign local platform coordinators who approve official accounts, support event publishers and maintain institution and department records.",
        "Begin with a small pilot between the University of Kigali and one partner institution, then expand using lessons from the pilot.",
        "Define measurable institutional indicators such as active collaborations, cross-institution event registrations, completed joint projects and response time to invitations.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.2.3 User-Adoption Recommendations", 3)
    for item in [
        "Provide short role-based training for students, lecturers, researchers, event organizers and administrators.",
        "Use academic champions in each institution to demonstrate valuable communities, projects and events during the pilot.",
        "Offer a simple onboarding checklist that helps users complete a profile, join a community, find a project and register for an event.",
        "Provide in-platform help, clear empty states and a support channel for users who experience access or notification problems.",
        "Collect structured feedback at regular intervals and publish visible improvements so that users see that their concerns influence development.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.2.4 Security Recommendations", 3)
    for item in [
        "Remove all default production credentials and fallback secrets; use a secret manager and rotate secrets regularly.",
        "Serve the platform only through HTTPS and WSS, restrict CORS to approved origins and place the service behind a hardened reverse proxy.",
        "Use secure session storage or carefully protected tokens, short token lifetimes, refresh-token rotation and multi-factor authentication for administrators.",
        "Scan uploads for malware, verify file content as well as extensions, randomize stored filenames and keep protected documents outside public static folders.",
        "Add rate limiting, login protection, security headers, detailed audit logs and alerts for sensitive administrative actions.",
        "Carry out vulnerability scanning, dependency review, penetration testing, backup restoration tests and incident-response exercises before launch.",
        "Adopt clear retention, consent, privacy, breach-response and account-deletion procedures aligned with applicable Rwandan data-protection requirements.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.2.5 Deployment Recommendations", 3)
    for item in [
        "Use separate development, testing and production environments with environment-specific configuration.",
        "Deploy the frontend, API, database and file storage as monitored services with daily automated backups and tested restoration.",
        "Use a continuous-integration pipeline that runs syntax checks, tests, dependency audits and the frontend production build before release.",
        "Add health monitoring, centralized logs, uptime alerts, database performance monitoring and Socket.IO connection metrics.",
        "Plan capacity for users, uploads and concurrent connections; use object storage and a shared real-time adapter when the platform is scaled horizontally.",
        "Release in phases: internal technical trial, limited institutional pilot, corrected release and then controlled multi-institution expansion.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.3 Areas for Further Research", 2)
    for item in [
        "A longitudinal multi-institution study comparing communication speed, collaboration activity and event participation before and after adoption.",
        "The performance and adoption of MQTT, COAR Notify or other event protocols alongside Socket.IO in Rwanda's higher-education environment.",
        "Real-time academic co-authoring using Operational Transformation or Conflict-free Replicated Data Types, including conflict, ownership and version-history management.",
        "The effect of low-bandwidth, offline-first and mobile designs on participation outside well-connected campuses.",
        "Trust, intellectual-property protection and willingness to share research across public and private institutions.",
        "Accessible multilingual interaction in English, Kinyarwanda and French for academic discovery and notifications.",
        "Privacy-preserving recommendation methods for matching users to projects, experts, communities and events.",
        "Governance and sustainability models for operating a national academic collaboration platform through institutional partnerships.",
    ]:
        add_bullet(doc, item)

    # Start the existing references on a new page.
    p_ref_break = doc.add_paragraph()
    p_ref_break.add_run().add_break(WD_BREAK.PAGE)

    # Move all newly created elements before the original REFERENCES heading.
    body = doc._element.body
    # python-docx appends new block elements immediately before the final sectPr.
    # Select that exact appended slice so original content is never moved or reordered.
    new_elements = list(body)[original_body_count - 1 : -1]
    ref_element = reference_paragraph._p
    for element in new_elements:
        body.remove(element)
        ref_element.addprevious(element)

    # Ask Word to refresh cached page-reference fields when the file is opened.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.core_properties.title = "Cross-Institutional Academic Collaboration and Real-Time Event Notification System"
    doc.core_properties.subject = "Completed Research Book with Chapters 4 and 5"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    result = build_document()
    print(result)
