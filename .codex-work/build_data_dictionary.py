from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"D:\Toon\My Doc\Classmate\Divine\System Data Dictionary - Section 4.6.4.docx"
)


def f(name, dtype, nullable, constraint, default, description):
    return {
        "name": name,
        "type": dtype,
        "nullable": nullable,
        "constraint": constraint,
        "default": default,
        "description": description,
    }


TABLES = [
    (
        "institutions",
        "Stores the higher-education and research institutions registered in the platform.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique institution identifier."),
            f("name", "VARCHAR(255)", "No", "UNIQUE", "None", "Official institution name."),
            f("type", "VARCHAR(100)", "No", "-", "None", "Institution category, such as University or Research Institute."),
            f("location", "VARCHAR(255)", "No", "-", "None", "Institution's physical location."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the institution record was created."),
        ],
    ),
    (
        "departments",
        "Stores departments and links each department to one institution.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique department identifier."),
            f("name", "VARCHAR(255)", "No", "UQ with institution_id", "None", "Department name; unique within its institution."),
            f("institution_id", "INT", "No", "FK -> institutions.id; ON DELETE CASCADE", "None", "Institution that owns the department."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the department was created."),
        ],
    ),
    (
        "users",
        "Stores user accounts, academic roles, affiliations and profile information.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique user identifier."),
            f("email", "VARCHAR(255)", "No", "UNIQUE", "None", "User's sign-in email address."),
            f("password_hash", "VARCHAR(255)", "No", "-", "None", "Secure hash of the user's password."),
            f("full_name", "VARCHAR(255)", "No", "-", "None", "User's full name."),
            f("role", "VARCHAR(50)", "No", "CHECK: student, lecturer, researcher, admin", "None", "Authorization role assigned to the user."),
            f("institution_id", "INT", "Yes", "FK -> institutions.id; ON DELETE SET NULL", "NULL", "Institution with which the user is affiliated."),
            f("department_id", "INT", "Yes", "FK -> departments.id; ON DELETE SET NULL", "NULL", "Department with which the user is affiliated."),
            f("bio", "TEXT", "Yes", "-", "NULL", "Short academic or professional biography."),
            f("avatar_url", "VARCHAR(255)", "Yes", "-", "NULL", "Path or URL of the user's profile image."),
            f("status", "VARCHAR(50)", "Yes", "CHECK: active, suspended", "'active'", "Current account status."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the account was created."),
        ],
    ),
    (
        "academic_communities",
        "Stores public, private or institution-specific academic communities.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique community identifier."),
            f("name", "VARCHAR(255)", "No", "UNIQUE", "None", "Community name."),
            f("description", "TEXT", "Yes", "-", "NULL", "Purpose and scope of the community."),
            f("category", "VARCHAR(100)", "No", "-", "None", "Academic category or subject area."),
            f("created_by", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User who created the community."),
            f("institution_id", "INT", "Yes", "FK -> institutions.id; ON DELETE SET NULL", "NULL", "Owning institution; NULL indicates a cross-institutional community."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the community was created."),
            f("privacy_type", "VARCHAR(50)", "Yes", "CHECK: public, private, institution", "'public'", "Controls who may discover or access the community."),
        ],
    ),
    (
        "community_members",
        "Links users to the academic communities they have joined.",
        [
            f("community_id", "INT", "No", "PK; FK -> academic_communities.id; ON DELETE CASCADE", "None", "Community membership identifier."),
            f("user_id", "INT", "No", "PK; FK -> users.id; ON DELETE CASCADE", "None", "User membership identifier."),
            f("joined_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the user joined."),
        ],
    ),
    (
        "community_invitations",
        "Stores invitations issued to users for private or restricted communities.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique invitation identifier."),
            f("community_id", "INT", "Yes", "FK -> academic_communities.id; ON DELETE CASCADE; UQ with user_id", "NULL", "Community to which the user is invited."),
            f("user_id", "INT", "Yes", "FK -> users.id; ON DELETE CASCADE; UQ with community_id", "NULL", "Invited user."),
            f("status", "VARCHAR(50)", "Yes", "CHECK: pending, accepted, rejected", "'pending'", "Current invitation decision."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the invitation was created."),
        ],
    ),
    (
        "posts",
        "Stores discussion posts published inside academic communities.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique post identifier."),
            f("title", "VARCHAR(255)", "No", "-", "None", "Post title."),
            f("content", "TEXT", "No", "-", "None", "Main post content."),
            f("user_id", "INT", "No", "FK -> users.id; ON DELETE CASCADE", "None", "User who authored the post."),
            f("community_id", "INT", "No", "FK -> academic_communities.id; ON DELETE CASCADE", "None", "Community in which the post was published."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the post was created."),
        ],
    ),
    (
        "comments",
        "Stores user comments attached to community posts.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique comment identifier."),
            f("post_id", "INT", "No", "FK -> posts.id; ON DELETE CASCADE", "None", "Post receiving the comment."),
            f("user_id", "INT", "No", "FK -> users.id; ON DELETE CASCADE", "None", "User who wrote the comment."),
            f("content", "TEXT", "No", "-", "None", "Comment text."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the comment was created."),
        ],
    ),
    (
        "likes",
        "Records which users have liked which community posts.",
        [
            f("post_id", "INT", "No", "PK; FK -> posts.id; ON DELETE CASCADE", "None", "Liked post identifier."),
            f("user_id", "INT", "No", "PK; FK -> users.id; ON DELETE CASCADE", "None", "User who liked the post."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the like was recorded."),
        ],
    ),
    (
        "projects",
        "Stores private academic collaboration projects and their access rules.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique project identifier."),
            f("title", "VARCHAR(255)", "No", "-", "None", "Project title."),
            f("description", "TEXT", "No", "-", "None", "Detailed project description."),
            f("created_by", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User who created the project."),
            f("status", "VARCHAR(50)", "Yes", "CHECK: planning, active, completed", "'planning'", "Current project lifecycle status."),
            f("requirements", "JSONB", "No", "CHECK: JSON value is an array", "'[]'::jsonb", "Structured list of skills or participation requirements."),
            f("privacy_type", "VARCHAR(20)", "No", "CHECK: private only", "'private'", "Project privacy mode."),
            f("access_scope", "VARCHAR(50)", "No", "CHECK: everyone, institution", "'everyone'", "Determines whether requests may come from everyone or only the institution."),
            f("institution_id", "INT", "Yes", "FK -> institutions.id; ON DELETE SET NULL", "NULL", "Institution used when access is institution-restricted."),
            f("updated_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time of the latest project update."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the project was created."),
        ],
    ),
    (
        "project_members",
        "Links users to projects and records their participation role.",
        [
            f("project_id", "INT", "No", "PK; FK -> projects.id; ON DELETE CASCADE", "None", "Project membership identifier."),
            f("user_id", "INT", "No", "PK; FK -> users.id; ON DELETE CASCADE", "None", "Participating user identifier."),
            f("role", "VARCHAR(50)", "Yes", "CHECK: lead, contributor, observer", "'contributor'", "User's responsibility within the project."),
            f("joined_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the user joined the project."),
        ],
    ),
    (
        "project_join_requests",
        "Stores user requests to join private academic projects.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique join-request identifier."),
            f("project_id", "INT", "No", "FK -> projects.id; ON DELETE CASCADE; UQ with user_id", "None", "Project the user wants to join."),
            f("user_id", "INT", "No", "FK -> users.id; ON DELETE CASCADE; UQ with project_id", "None", "User requesting project membership."),
            f("status", "VARCHAR(50)", "No", "CHECK: pending, accepted, rejected", "'pending'", "Current decision on the request."),
            f("requested_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the request was submitted."),
            f("responded_at", "TIMESTAMPTZ", "Yes", "-", "NULL", "Date and time when the request was decided."),
            f("responded_by", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User who accepted or rejected the request."),
        ],
    ),
    (
        "project_files",
        "Stores metadata for files uploaded to collaboration projects.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique project-file identifier."),
            f("project_id", "INT", "No", "FK -> projects.id; ON DELETE CASCADE", "None", "Project that owns the file."),
            f("filename", "VARCHAR(255)", "No", "-", "None", "Original or displayed file name."),
            f("filepath", "VARCHAR(255)", "No", "-", "None", "Server-side storage path."),
            f("mime_type", "VARCHAR(255)", "Yes", "-", "NULL", "File's media type."),
            f("file_size", "BIGINT", "Yes", "-", "NULL", "File size in bytes."),
            f("uploaded_by", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User who uploaded the file."),
            f("uploaded_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the file was uploaded."),
        ],
    ),
    (
        "events",
        "Stores academic events, schedules, venues and participation capacity.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique event identifier."),
            f("title", "VARCHAR(255)", "No", "-", "None", "Event title."),
            f("description", "TEXT", "Yes", "-", "NULL", "Detailed event information."),
            f("event_date", "TIMESTAMPTZ", "No", "-", "None", "Scheduled date and time of the event."),
            f("location", "VARCHAR(255)", "No", "-", "None", "Physical venue or location description."),
            f("meeting_link", "VARCHAR(1000)", "Yes", "-", "NULL", "Online meeting URL, when applicable."),
            f("organizer_id", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User responsible for organizing the event."),
            f("institution_id", "INT", "Yes", "FK -> institutions.id; ON DELETE SET NULL", "NULL", "Host institution; NULL means the event is open across institutions."),
            f("capacity", "INT", "Yes", "-", "100", "Maximum planned number of participants."),
            f("updated_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time of the latest event update."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the event was created."),
        ],
    ),
    (
        "event_registrations",
        "Links users to the academic events for which they have registered.",
        [
            f("event_id", "INT", "No", "PK; FK -> events.id; ON DELETE CASCADE", "None", "Registered event identifier."),
            f("user_id", "INT", "No", "PK; FK -> users.id; ON DELETE CASCADE", "None", "Registered user identifier."),
            f("registered_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when registration occurred."),
        ],
    ),
    (
        "news",
        "Stores public academic news and announcements.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique news item identifier."),
            f("title", "VARCHAR(255)", "No", "-", "None", "News headline."),
            f("description", "TEXT", "No", "-", "None", "News content or summary."),
            f("category", "VARCHAR(50)", "No", "-", "'Other'", "News category."),
            f("feature_image", "VARCHAR(500)", "No", "-", "None", "Path or URL of the main news image."),
            f("external_link", "VARCHAR(1000)", "Yes", "-", "NULL", "Optional link to an external source or related page."),
            f("created_by", "INT", "Yes", "FK -> users.id; ON DELETE SET NULL", "NULL", "User who published the news item."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the item was created."),
            f("updated_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time of the latest news update."),
        ],
    ),
    (
        "news_documents",
        "Stores document attachments associated with public news items.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique news-document identifier."),
            f("news_id", "INT", "No", "FK -> news.id; ON DELETE CASCADE", "None", "News item that owns the document."),
            f("filename", "VARCHAR(255)", "No", "-", "None", "Original or displayed document name."),
            f("filepath", "VARCHAR(500)", "No", "-", "None", "Server-side storage path."),
            f("mime_type", "VARCHAR(150)", "Yes", "-", "NULL", "Document media type."),
            f("file_size", "BIGINT", "Yes", "-", "NULL", "Document size in bytes."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the document was attached."),
        ],
    ),
    (
        "chat_rooms",
        "Stores one-to-one and group chat rooms.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique chat-room identifier."),
            f("name", "VARCHAR(255)", "Yes", "-", "NULL", "Room name; NULL is used for a one-to-one conversation."),
            f("is_group", "BOOLEAN", "Yes", "-", "FALSE", "Indicates whether the room is a group chat."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the room was created."),
        ],
    ),
    (
        "chat_members",
        "Links users to the chat rooms in which they participate.",
        [
            f("room_id", "INT", "No", "PK; FK -> chat_rooms.id; ON DELETE CASCADE", "None", "Chat-room membership identifier."),
            f("user_id", "INT", "No", "PK; FK -> users.id; ON DELETE CASCADE", "None", "Participating user identifier."),
            f("joined_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the user joined the room."),
        ],
    ),
    (
        "chat_messages",
        "Stores messages sent within chat rooms.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique message identifier."),
            f("room_id", "INT", "No", "FK -> chat_rooms.id; ON DELETE CASCADE", "None", "Chat room containing the message."),
            f("sender_id", "INT", "No", "FK -> users.id; ON DELETE CASCADE", "None", "User who sent the message."),
            f("message", "TEXT", "No", "-", "None", "Message content."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the message was sent."),
        ],
    ),
    (
        "notifications",
        "Stores persistent in-application notifications delivered to individual users.",
        [
            f("id", "SERIAL", "No", "PK", "Auto-generated", "Unique notification identifier."),
            f("user_id", "INT", "No", "FK -> users.id; ON DELETE CASCADE", "None", "User receiving the notification."),
            f("title", "VARCHAR(255)", "No", "-", "None", "Short notification heading."),
            f("content", "TEXT", "No", "-", "None", "Notification message."),
            f("type", "VARCHAR(50)", "No", "Application values: chat, event, project, community, system", "None", "Notification category used by the interface."),
            f("is_read", "BOOLEAN", "Yes", "-", "FALSE", "Indicates whether the user has opened or acknowledged the notification."),
            f("link", "VARCHAR(255)", "Yes", "-", "NULL", "Optional internal destination related to the notification."),
            f("created_at", "TIMESTAMPTZ", "Yes", "-", "CURRENT_TIMESTAMP", "Date and time when the notification was created."),
        ],
    ),
]


def set_cell_text(cell, text, *, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(cell, top=50, start=80, bottom=50, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "80")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def set_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.49)
    sec.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    set_font(normal, "Times New Roman", 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5

    h1 = doc.styles["Heading 1"]
    set_font(h1, "Times New Roman", 16, True, "000000")
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_font(h2, "Times New Roman", 12, True, "000000")
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    set_font(caption, "Times New Roman", 10, True, "000000")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    footer = sec.footer
    add_page_number(footer.paragraphs[0])
    for r in footer.paragraphs[0].runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("4.6.4 Data Dictionary")

    intro = doc.add_paragraph()
    intro.add_run(
        "The following data dictionary documents the twenty-one tables implemented in the "
        "PostgreSQL database of the Cross-Institutional Academic Collaboration and Real-Time "
        "Event Notification System. It is based directly on the supplied system schema. "
        "Each table identifies the stored fields, PostgreSQL data types, null rules, keys, "
        "constraints, default values and the purpose of each field."
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(8)
    note.paragraph_format.line_spacing = 1.15
    r = note.add_run(
        "Notation: PK = Primary Key; FK = Foreign Key; UQ = Unique constraint; "
        "TIMESTAMPTZ = timestamp with time zone. A column marked “Yes” under Nullable "
        "accepts NULL according to the implemented SQL schema."
    )
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)

    widths = [1320, 1420, 620, 2100, 1240, 2660]
    headers = ["Field", "Data type", "Null?", "Key / constraint", "Default", "Description"]

    # Deliberate group boundaries keep captions away from page edges and make
    # the standalone tables easier to copy into the main research book.
    page_break_before = {3, 4, 6, 9, 12, 14, 15, 17, 20}

    for idx, (table_name, purpose, fields) in enumerate(TABLES, start=1):
        cap = doc.add_paragraph(style="Caption")
        if idx in page_break_before:
            cap.paragraph_format.page_break_before = True
        cap.add_run(f"Table 4.6.4-{idx}: Data Dictionary for {table_name}")

        purpose_p = doc.add_paragraph()
        purpose_p.paragraph_format.space_before = Pt(0)
        purpose_p.paragraph_format.space_after = Pt(4)
        purpose_p.paragraph_format.line_spacing = 1.0
        purpose_p.paragraph_format.keep_with_next = True
        label = purpose_p.add_run("Purpose: ")
        label.bold = True
        label.font.name = "Times New Roman"
        label.font.size = Pt(9)
        text = purpose_p.add_run(purpose)
        text.font.name = "Times New Roman"
        text.font.size = Pt(9)

        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for i, header in enumerate(headers):
            set_cell_text(table.rows[0].cells[i], header, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_margins(table.rows[0].cells[i])
        set_repeat_table_header(table.rows[0])
        prevent_row_split(table.rows[0])

        for field in fields:
            row = table.add_row()
            values = [
                field["name"],
                field["type"],
                field["nullable"],
                field["constraint"],
                field["default"],
                field["description"],
            ]
            for i, value in enumerate(values):
                align = WD_ALIGN_PARAGRAPH.CENTER if i in (1, 2, 4) else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(row.cells[i], value, size=8.3, align=align)
                set_cell_margins(row.cells[i])
            prevent_row_split(row)

        set_table_geometry(table, widths)
        set_table_borders(table)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = "System Data Dictionary - Section 4.6.4"
    doc.core_properties.subject = "Cross-Institutional Academic Collaboration and Real-Time Event Notification System"
    doc.core_properties.author = "Researcher"
    doc.core_properties.keywords = "data dictionary, PostgreSQL, academic collaboration"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Tables documented: {len(TABLES)}")


if __name__ == "__main__":
    build()
